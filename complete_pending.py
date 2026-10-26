#!/usr/bin/env python3
"""
complete_pending.py —— 独立 LLM 补全脚本（替代上游 Hermes agent 的补全步骤）

对 papers_record.xlsx 中缺失 affiliations / summary_cn 的论文：
  读取 PDF 前两页文本 -> 调用大模型（Claude/Anthropic 兼容协议）->
  写回 Excel -> 维护 merged_llm_results.json -> 调 monitor.py --sync-pending-state

配置从仓库外读取：<repo 的上级目录>/pipeline_config.json
  { "baseURL": "https://api.kimi.com/coding", "model": "k3[1m]", "apiKey": "sk-..." }

用法：
  python complete_pending.py            # 补全全部待处理论文
  python complete_pending.py --limit 5  # 只处理前 5 篇（试跑）
"""

from __future__ import annotations

import json
import re
import subprocess
import sys
import time
import urllib.request
import urllib.error
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
import pdfplumber

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR.parent / "data"
CONFIG_PATH = BASE_DIR.parent / "pipeline_config.json"
EXCEL_PATH = DATA_DIR / "papers_record.xlsx"
PAPERS_DIR = DATA_DIR / "papers"
MERGED_PATH = DATA_DIR / "merged_llm_results.json"
MONITOR = BASE_DIR / "monitor.py"

PDF_MAX_CHARS = 6000
HTTP_TIMEOUT = 180
RETRIES = 2
SAVE_EVERY = 5


def load_config() -> dict:
    if not CONFIG_PATH.exists():
        raise SystemExit(f"[ERROR] 未找到配置文件：{CONFIG_PATH}（含 baseURL/model/apiKey）")
    cfg = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    for k in ("baseURL", "model", "apiKey"):
        if not cfg.get(k):
            raise SystemExit(f"[ERROR] 配置缺少 {k}")
    return cfg


def call_llm(cfg: dict, prompt: str) -> str:
    root = cfg["baseURL"].rstrip().rstrip("/")
    if root.endswith("/v1"):
        root = root[: -len("/v1")]
    url = root + "/v1/messages"
    body = json.dumps({
        "model": cfg["model"],
        "max_tokens": 1200,
        "temperature": 0.2,
        "messages": [{"role": "user", "content": prompt}],
    }).encode("utf-8")
    req = urllib.request.Request(url, data=body, method="POST", headers={
        "Content-Type": "application/json",
        "x-api-key": cfg["apiKey"],
        "anthropic-version": "2023-06-01",
    })
    with urllib.request.urlopen(req, timeout=HTTP_TIMEOUT) as resp:
        data = json.loads(resp.read().decode("utf-8"))
    return "".join(b.get("text", "") for b in data.get("content", []) if isinstance(b, dict))


def extract_json(text: str) -> dict | None:
    t = text.strip()
    m = re.search(r"\{[\s\S]*\}", t)
    if not m:
        return None
    try:
        j = json.loads(m.group(0))
        if isinstance(j, dict):
            return j
    except json.JSONDecodeError:
        pass
    return None


def pdf_text(arxiv_id: str) -> str:
    pdf = PAPERS_DIR / f"{arxiv_id}.pdf"
    if not pdf.exists():
        return ""
    parts = []
    try:
        with pdfplumber.open(pdf) as doc:
            for page in doc.pages[:2]:
                parts.append(page.extract_text() or "")
    except Exception as e:  # noqa: BLE001 损坏 PDF 不应中断整批
        print(f"    [WARN] PDF 读取失败 {arxiv_id}: {e}")
    return "\n".join(parts)[:PDF_MAX_CHARS]


def build_prompt(paper: dict, pdf: str) -> str:
    return (
        "你是论文整理助手。根据下面的论文信息与 PDF 前两页文本，完成两项任务：\n"
        "1. affiliations：提取全部作者单位（按 PDF 署名顺序，去重，用\"; \"分隔，"
        "保留机构原始英文名称，如 \"Tsinghua University; JD.com\"；若确实无法确定，填\"未找到单位信息\"）。\n"
        "2. summary_cn：写 90-150 字中文总结，覆盖：解决什么问题、方法核心思路、主要结果（保留关键数字）。"
        "直接陈述，不要以\"这篇论文\"开头，不要分点，不要出现英文整句。\n"
        "严格只输出 JSON（无代码块标记、无解释）："
        "{\"affiliations\":\"...\",\"summary_cn\":\"...\"}\n\n"
        f"标题: {paper.get('title', '')}\n"
        f"作者: {paper.get('authors', '')}\n"
        f"摘要: {str(paper.get('abstract', ''))[:800]}\n\n"
        f"PDF 前两页文本:\n{pdf if pdf else '（PDF 文本不可用，请基于摘要推断，单位填\"未找到单位信息\"）'}"
    )


def load_incomplete() -> list[dict]:
    wb = load_workbook(EXCEL_PATH, read_only=True)
    ws = wb["Papers"]
    headers = [str(h) if h is not None else "" for h in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    idx = {name: i for i, name in enumerate(headers)}
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        def g(col: str) -> str:
            v = row[idx[col]] if idx.get(col) is not None and idx[col] < len(row) else None
            return str(v).strip() if v is not None else ""
        arxiv_id = g("arxiv_id")
        if not arxiv_id:
            continue
        if not g("summary_cn") or not g("affiliations"):
            out.append({
                "arxiv_id": arxiv_id,
                "title": g("title"),
                "authors": g("authors"),
                "abstract": g("abstract"),
                "row": None,
            })
    wb.close()
    return out


def update_excel(results: list[dict]) -> int:
    """把补全结果写回 Excel（按 arxiv_id 定位行），返回更新条数。"""
    if not results:
        return 0
    wb = load_workbook(EXCEL_PATH)
    ws = wb["Papers"]
    row_by_id = {}
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v:
            row_by_id[str(v).strip()] = r
    updated = 0
    for res in results:
        r = row_by_id.get(res["arxiv_id"])
        if not r:
            continue
        if res.get("affiliations"):
            ws.cell(row=r, column=4).value = res["affiliations"]
        if res.get("summary_cn"):
            ws.cell(row=r, column=8).value = res["summary_cn"]
        updated += 1
    wb.save(EXCEL_PATH)
    return updated


def merge_results_file(results: list[dict]) -> None:
    merged = {}
    if MERGED_PATH.exists():
        try:
            merged = json.loads(MERGED_PATH.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            merged = {}
    for res in results:
        merged[res["arxiv_id"]] = {
            "affiliations": res.get("affiliations", ""),
            "summary_cn": res.get("summary_cn", ""),
            "chinese_count": len(res.get("summary_cn", "") or ""),
        }
    MERGED_PATH.write_text(json.dumps(merged, ensure_ascii=False, indent=2), encoding="utf-8")


def sync_pending() -> None:
    subprocess.run([sys.executable, str(MONITOR), "--sync-pending-state"], cwd=BASE_DIR, check=False)


def main() -> None:
    limit = 0
    if "--limit" in sys.argv:
        limit = int(sys.argv[sys.argv.index("--limit") + 1])

    cfg = load_config()
    todo = load_incomplete()
    print(f"[INFO] 待补全论文：{len(todo)} 篇 | 模型：{cfg['model']}")
    if limit:
        todo = todo[:limit]
        print(f"[INFO] --limit 生效，本轮处理 {len(todo)} 篇")
    if not todo:
        sync_pending()
        print("[OK] 无待补全论文")
        return

    done, failed = [], []
    for i, paper in enumerate(todo, 1):
        arxiv_id = paper["arxiv_id"]
        print(f"[{i}/{len(todo)}] {arxiv_id} | {paper['title'][:60]}")
        pdf = pdf_text(arxiv_id)
        if not pdf and not (PAPERS_DIR / f"{arxiv_id}.pdf").exists():
            # PDF 永久缺失（下载失败/损坏）：不再跳过——否则该论文永远 pending，卡死发布门禁。
            # 改用摘要兜底生成 summary_cn，单位标注"未找到单位信息"。
            print("    [INFO] 本地无 PDF，用摘要兜底补全")
        res = None
        for attempt in range(1, RETRIES + 1):
            try:
                raw = call_llm(cfg, build_prompt(paper, pdf))
                j = extract_json(raw)
                if j and j.get("summary_cn"):
                    res = {
                        "arxiv_id": arxiv_id,
                        "affiliations": str(j.get("affiliations", "")).strip(),
                        "summary_cn": str(j.get("summary_cn", "")).strip(),
                    }
                    break
                print(f"    [WARN] 第 {attempt} 次返回无法解析：{raw[:80]}")
            except (urllib.error.URLError, urllib.error.HTTPError, TimeoutError, OSError) as e:
                print(f"    [WARN] 第 {attempt} 次请求失败：{e}")
                time.sleep(3 * attempt)
        if res:
            done.append(res)
            print(f"    [OK] 单位: {res['affiliations'][:60]} | 摘要 {len(res['summary_cn'])} 字")
        else:
            failed.append(arxiv_id)
            print("    [FAIL] 放弃本篇，留待下轮")

        if len(done) and len(done) % SAVE_EVERY == 0:
            n = update_excel(done[-SAVE_EVERY:])
            merge_results_file(done[-SAVE_EVERY:])
            print(f"    [SAVE] 已写回 Excel {n} 条")

    if len(done) % SAVE_EVERY:
        n = update_excel(done[len(done) - (len(done) % SAVE_EVERY):])
        merge_results_file(done[len(done) - (len(done) % SAVE_EVERY):])
        print(f"[SAVE] 已写回 Excel {n} 条")

    sync_pending()
    tail_msg = ('（失败: ' + ', '.join(failed[:5]) + ('…' if len(failed) > 5 else '') + '）') if failed else ''
    print(f"[DONE] {datetime.now().isoformat()} 成功 {len(done)} / 失败 {len(failed)}" + tail_msg)


if __name__ == "__main__":
    main()
